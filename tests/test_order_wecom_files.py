import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from test_store_order_auth_gate import load_inventory_module


WEBHOOK = "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=test-only"


def excel_file(tmp_path):
    path = tmp_path / "银泰城店_日配订货.xlsx"
    book = Workbook()
    book.active.append(["门店", "商品", "数量"])
    book.active.append(["银泰城店", "冷冻牛排", 6])
    book.save(path)
    book.close()
    return path


def response(payload):
    return io.BytesIO(json.dumps(payload).encode())


def test_order_notification_sends_original_excel_after_order_text(tmp_path, monkeypatch):
    module = load_inventory_module()
    path = excel_file(tmp_path)
    original = path.read_bytes()
    calls = []
    logs = []

    def send(req, **kwargs):
        calls.append(req)
        if "upload_media" in req.full_url:
            return response({"errcode": 0, "media_id": "uploaded-order"})
        return response({"errcode": 0, "errmsg": "ok"})

    monkeypatch.setenv("ORDER_NOTIFY_TYPE", "wecom")
    monkeypatch.setenv("ORDER_NOTIFY_WEBHOOK", WEBHOOK)
    monkeypatch.setattr(module.url_request, "urlopen", send)
    monkeypatch.setattr(module, "_write_order_notify_log", lambda *args: logs.append(args))
    module._notify_order_submit({"file": str(path), "items": [
        {"store_name": "银泰城店", "product_name": "冷冻牛排", "quantity": 6, "unit": "袋"},
    ]}, path.name, "http://example.test/order.xlsx")

    assert len(calls) == 3
    text = json.loads(calls[0].data)["text"]["content"]
    assert "门店：银泰城店" in text and "冷冻牛排 6袋" in text
    assert "下载：http://example.test/order.xlsx" in text
    assert "upload_media" in calls[1].full_url
    # Extract exactly the uploaded attachment and compare all original bytes.
    uploaded = calls[1].data.split(b"\r\n\r\n", 1)[1].rsplit(b"\r\n--", 1)[0]
    assert uploaded == original
    book = load_workbook(io.BytesIO(uploaded))
    assert book.active["A2"].value == "银泰城店"
    assert book.active["B2"].value == "冷冻牛排"
    assert book.active["C2"].value == 6
    book.close()
    assert json.loads(calls[2].data) == {"msgtype": "file", "file": {"media_id": "uploaded-order"}}
    assert logs[-1][0] == "file-sent"
    assert path.read_bytes() == original


@pytest.mark.parametrize("reply", [{"errcode": 40058}, {"errcode": 0}])
def test_upload_failure_never_sends_a_fake_attachment(tmp_path, monkeypatch, reply):
    helper = load_inventory_module().wecom_files
    calls = []

    def send(req, **kwargs):
        calls.append(req)
        return response(reply)

    monkeypatch.setattr(helper.request, "urlopen", send)
    with pytest.raises(ValueError):
        helper.send_excel(WEBHOOK, excel_file(tmp_path))
    assert len(calls) == 1


def test_non_excel_content_is_rejected_before_upload(tmp_path, monkeypatch):
    helper = load_inventory_module().wecom_files
    path = tmp_path / "假的Excel.xlsx"
    path.write_text("这只是一个链接，不是Excel", encoding="utf-8")
    monkeypatch.setattr(helper.request, "urlopen", lambda *a, **k: pytest.fail("must not upload"))
    with pytest.raises(Exception):
        helper.send_excel(WEBHOOK, path)


def test_file_api_error_is_not_reported_as_success(tmp_path, monkeypatch):
    helper = load_inventory_module().wecom_files
    replies = iter([{"errcode": 0, "media_id": "test"}, {"errcode": 45009}])
    monkeypatch.setattr(helper.request, "urlopen", lambda *a, **k: response(next(replies)))
    with pytest.raises(ValueError, match="45009"):
        helper.send_excel(WEBHOOK, excel_file(tmp_path))
