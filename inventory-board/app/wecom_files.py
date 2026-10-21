"""Send an existing Excel workbook through the same WeCom group webhook."""
from __future__ import annotations

import hashlib
import io
import json
import uuid
from pathlib import Path
from urllib import request
from urllib.parse import parse_qs, urlencode, urlsplit

from openpyxl import load_workbook


def _post_json(url: str, payload: dict) -> dict:
    req = request.Request(
        url, data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method="POST",
    )
    with request.urlopen(req, timeout=20) as response:
        result = json.load(response)
    _require_success(result)
    return result


def _require_success(result: dict) -> None:
    if result.get("errcode") != 0:
        # Never include the webhook URL (or key) in error messages.
        raise ValueError(f"WeCom API error: {result.get('errcode')}")


def send_excel(webhook: str, path: Path) -> dict:
    """Validate and upload the original bytes, then send a real file message.

    No automatic retries: an ambiguous send timeout must not duplicate orders.
    """
    parsed = urlsplit(webhook)
    keys = parse_qs(parsed.query).get("key", [])
    if (parsed.scheme != "https" or parsed.hostname != "qyapi.weixin.qq.com"
            or parsed.path != "/cgi-bin/webhook/send" or len(keys) != 1):
        raise ValueError("Invalid WeCom group webhook")
    if path.suffix.lower() != ".xlsx" or any(c in path.name for c in '\r\n"'):
        raise ValueError("Expected an Excel .xlsx filename")
    data = path.read_bytes()
    if not 5 <= len(data) <= 20 * 1024 * 1024:
        raise ValueError("Excel attachment size is outside WeCom limits")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if not workbook.sheetnames:
            raise ValueError("Excel workbook has no worksheets")
        # Read every sheet so malformed XML does not pass a header-only check.
        for sheet in workbook:
            for row in sheet.iter_rows(values_only=True):
                pass
    finally:
        workbook.close()
    boundary = "----WeComExcel" + uuid.uuid4().hex
    body = (
        f'--{boundary}\r\nContent-Disposition: form-data; name="media"; '
        f'filename="{path.name}"; filelength={len(data)}\r\n'
        'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'
    ).encode("utf-8") + data + f"\r\n--{boundary}--\r\n".encode()
    upload_url = "https://qyapi.weixin.qq.com/cgi-bin/webhook/upload_media?" + urlencode(
        {"key": keys[0], "type": "file"},
    )
    req = request.Request(upload_url, data=body, method="POST", headers={
        "Content-Type": f"multipart/form-data; boundary={boundary}",
    })
    with request.urlopen(req, timeout=30) as response:
        uploaded = json.load(response)
    _require_success(uploaded)
    if not uploaded.get("media_id"):
        raise ValueError("WeCom upload returned no media_id")
    _post_json(webhook, {"msgtype": "file", "file": {"media_id": uploaded["media_id"]}})
    return {"filename": path.name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
